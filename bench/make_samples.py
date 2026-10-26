import io
import os
import random
import subprocess
import sys
from datetime import date, timedelta

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
random.seed(42)


def ktp():
    from PIL import Image, ImageDraw, ImageFilter
    w, h = 1600, 1010
    img = Image.new("RGB", (w, h), (154, 196, 226))
    d = ImageDraw.Draw(img)
    for y in range(0, h, 4):
        d.line([(0, y), (w, y + 40)], fill=(150 + (y % 13), 192, 222), width=2)
    d.rectangle([40, 40, w - 40, h - 40], outline=(40, 70, 120), width=6)
    d.text((70, 70), "SPECIMEN - NOT A REAL DOCUMENT - FICTIONAL TEST DATA", fill=(170, 20, 20))
    rows = ["NIK : 3578 0000 0000 0000", "Name : EXAMPLE PATIENT", "Place/Date of birth : SURABAYA, 01-01-1990",
            "Address : JL. CONTOH NO. 1", "Blood type : O", "Valid until : LIFETIME"]
    for i, r in enumerate(rows):
        d.text((90, 200 + i * 90), r, fill=(20, 20, 40))
    d.rectangle([1150, 200, 1480, 640], fill=(210, 210, 215), outline=(60, 60, 60), width=3)
    d.ellipse([1240, 260, 1390, 420], fill=(150, 150, 160))
    d.rectangle([1210, 430, 1420, 620], fill=(150, 150, 160))
    noise = Image.effect_noise((w, h), 28).convert("RGB")
    img = Image.blend(img, noise, 0.10).filter(ImageFilter.GaussianBlur(0.6))
    img.save(os.path.join(OUT, "ktp_specimen.jpg"), quality=90)


def xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "BP & sugar"
    ws.append(["Date", "Systolic", "Diastolic", "Pulse", "Fasting glucose (mg/dL)", "Note"])
    d0 = date(2025, 10, 1)
    for i in range(365):
        ws.append([d0 + timedelta(days=i), random.randint(118, 152), random.randint(74, 96), random.randint(62, 90),
                   random.randint(88, 140), random.choice(["", "", "mild dizziness", "missed medication", ""])])
    wb.save(os.path.join(OUT, "bp_sugar_log.xlsx"))


def docx():
    from docx import Document
    doc = Document()
    doc.add_heading("DISCHARGE SUMMARY (SPECIMEN - FICTIONAL)", 1)
    doc.add_paragraph("Name: Example Patient    MRN: 000000    Unit: Neurology")
    for title in ("History", "Physical examination", "Diagnosis", "Treatment", "Follow-up plan"):
        doc.add_heading(title, 2)
        for _ in range(4):
            doc.add_paragraph(" ".join(random.choice(
                ["patient", "seizure", "recurrent", "duration", "minutes", "without", "fever", "history", "family",
                 "dose", "follow-up", "EEG", "normal", "therapy", "continue", "month", "evaluation"]) for _ in range(40)))
    doc.save(os.path.join(OUT, "discharge_summary.docx"))


def pdf():
    from PIL import Image
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    buf = os.path.join(OUT, "lab_result.pdf")
    c = canvas.Canvas(buf, pagesize=A4)
    for page in range(3):
        scan = Image.effect_noise((1240, 1754), 22).convert("L").point(lambda v: 190 + v // 4)
        bio = io.BytesIO()
        scan.save(bio, format="JPEG", quality=88)
        bio.seek(0)
        c.drawImage(ImageReader(bio), 0, 0, *A4)
        c.setFont("Helvetica-Bold", 14)
        c.drawString(60, 790, f"LABORATORY RESULTS (SPECIMEN, FICTIONAL) - page {page + 1}/3")
        c.setFont("Helvetica", 10)
        y = 740
        for name in ("Haemoglobin", "Leukocytes", "Platelets", "Fasting glucose", "HbA1c", "Creatinine", "Urea",
                     "AST", "ALT", "Total cholesterol", "Sodium", "Potassium"):
            c.drawString(60, y, f"{name:<20} {random.uniform(1, 200):7.1f}")
            y -= 22
        c.showPage()
    c.save()


def video(name, seconds, bitrate):
    path = os.path.join(OUT, name)
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"testsrc2=size=1280x720:rate=30:d={seconds}",
           "-vf", "noise=alls=18:allf=t", "-c:v", "libx264", "-preset", "veryfast", "-b:v", bitrate,
           "-maxrate", bitrate, "-bufsize", "8M", "-pix_fmt", "yuv420p", "-movflags", "+faststart", path]
    subprocess.run(cmd, check=True)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    ktp(); xlsx(); docx(); pdf()
    if "--no-video" not in sys.argv:
        video("symptom_video_30s.mp4", 30, "4M")
        video("symptom_video_3min.mp4", 180, "4M")
    for f in sorted(os.listdir(OUT)):
        print(f"{f:28} {os.path.getsize(os.path.join(OUT, f)) / 1e6:8.2f} MB")
